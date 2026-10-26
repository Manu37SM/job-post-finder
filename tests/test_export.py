import csv as csv_module

import openpyxl

from export.csv import export_csv
from export.excel import export_excel
from export.text import export_text
from tests.factories import make_record


def test_export_excel(tmp_path):
    records = [
        make_record(post_url="https://linkedin.com/posts/1"),
        make_record(
            post_url="https://linkedin.com/posts/2",
            company="XYZ Corp",
            primary_email="careers@xyz.com",
        ),
    ]
    path = export_excel(records, tmp_path / "out.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "post_url"
    assert "email" in header
    assert "role_confidence" in header

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    email_col = header.index("email")
    assert rows[0][email_col] == "jobs@example.com"
    assert rows[1][email_col] == "careers@xyz.com"


def test_export_excel_empty_list_does_not_raise(tmp_path):
    path = export_excel([], tmp_path / "empty.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.active.max_row == 1  # header only


def test_export_csv(tmp_path):
    path = export_csv([make_record()], tmp_path / "out.csv")
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv_module.DictReader(f))
    assert len(rows) == 1
    assert rows[0]["email"] == "jobs@example.com"
    assert rows[0]["job_role"] == "Java Developer"
    assert rows[0]["work_mode"] == "ONSITE"


def test_export_csv_empty_list_writes_header_only(tmp_path):
    path = export_csv([], tmp_path / "empty.csv")
    with open(path, newline="", encoding="utf-8") as f:
        rows = list(csv_module.DictReader(f))
    assert rows == []


def test_export_text(tmp_path):
    path = export_text([make_record()], tmp_path / "out.txt")
    content = path.read_text(encoding="utf-8")
    assert "Java Developer — Mumbai" in content
    assert "jobs@example.com" in content
    assert "LinkedIn Post:" in content
    # spec's own example block has 3 separator lines per record: before the
    # heading, after the heading, and closing the record.
    assert content.count("=" * 50) == 3


def test_export_text_unknown_fields_render_as_unknown(tmp_path):
    record = make_record(company=None, job_role=None, job_location=None, primary_email=None)
    path = export_text([record], tmp_path / "out.txt")
    content = path.read_text(encoding="utf-8")
    assert "UNKNOWN — UNKNOWN" in content


def test_export_text_empty_list_writes_empty_file(tmp_path):
    path = export_text([], tmp_path / "empty.txt")
    assert path.read_text(encoding="utf-8") == ""


def test_excel_export_survives_control_characters_from_ocr(tmp_path):
    """Regression (v11): one \\x0c in OCR text raised IllegalCharacterError for the whole export."""
    from export.excel import export_excel
    from tests.factories import make_record

    rec = make_record(post_text="Java\x0bDeveloper\x00 hiring")
    path = export_excel([rec], tmp_path / "x.xlsx")
    from openpyxl import load_workbook

    ws = load_workbook(path).active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert any(isinstance(v, str) and "JavaDeveloper hiring" in v for v in values)


def test_csv_has_bom_for_excel(tmp_path):
    from export.csv import export_csv
    from tests.factories import make_record

    path = export_csv([make_record(author_name="Zoë")], tmp_path / "x.csv")
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
