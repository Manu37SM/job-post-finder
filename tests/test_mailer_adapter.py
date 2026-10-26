import openpyxl

from core.models import RecordStatus
from export.mailer_adapter import export_approved_for_mailer, to_mailer_rows
from tests.factories import make_record


def test_only_approved_records_exported():
    records = [
        make_record(post_url="https://linkedin.com/posts/1", status=RecordStatus.APPROVED),
        make_record(post_url="https://linkedin.com/posts/2", status=RecordStatus.NEW),
        make_record(post_url="https://linkedin.com/posts/3", status=RecordStatus.REJECTED),
        make_record(post_url="https://linkedin.com/posts/4", status=RecordStatus.SKIPPED),
    ]
    rows = to_mailer_rows(records)
    assert len(rows) == 1
    assert rows[0]["Email"] == "jobs@example.com"


def test_approved_record_missing_email_excluded():
    records = [
        make_record(status=RecordStatus.APPROVED, primary_email=None, all_emails=[]),
    ]
    assert to_mailer_rows(records) == []


def test_row_shape_matches_mass_mailer_schema():
    records = [
        make_record(
            status=RecordStatus.APPROVED,
            author_name="Jane Recruiter",
            primary_email="jane@abc.com",
            company="ABC Technologies",
            job_role="Java Developer",
        )
    ]
    rows = to_mailer_rows(records)
    assert rows[0] == {
        "Name": "Jane Recruiter",
        "Email": "jane@abc.com",
        "Company": "ABC Technologies",
        "Title": "Java Developer",
    }


def test_missing_optional_fields_become_empty_string_not_unknown():
    # mass-mailer's parser does str(row["Company"]).strip() - "UNKNOWN" would
    # be sent to a recruiter as their own company name, so these must be "".
    records = [
        make_record(
            status=RecordStatus.APPROVED,
            company=None,
            job_role=None,
            primary_email="jobs@example.com",
        )
    ]
    rows = to_mailer_rows(records)
    assert rows[0]["Company"] == ""
    assert rows[0]["Title"] == ""


def test_export_writes_correct_headers_and_rows(tmp_path):
    records = [
        make_record(status=RecordStatus.APPROVED),
        make_record(post_url="https://linkedin.com/posts/2", status=RecordStatus.NEW),
    ]
    path = export_approved_for_mailer(records, tmp_path / "approved.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["Name", "Email", "Company", "Title"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1  # only the APPROVED one
    assert rows[0][1] == "jobs@example.com"


def test_export_empty_approved_list_writes_header_only(tmp_path):
    path = export_approved_for_mailer([], tmp_path / "approved.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.active.max_row == 1


def test_same_email_approved_twice_is_exported_once():
    from core.models import RecordStatus
    from export.mailer_adapter import to_mailer_rows
    from tests.factories import make_record

    a = make_record(post_url="u1", primary_email="HR@Acme.com", status=RecordStatus.APPROVED)
    b = make_record(post_url="u2", primary_email="hr@acme.com", status=RecordStatus.APPROVED)
    assert len(to_mailer_rows([a, b])) == 1
