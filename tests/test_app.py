import csv
import json

import openpyxl

from app import run_pipeline
from config import Config
from linkedin.fixture_source import discover_posts
from storage.checkpoint import read_checkpoint
from storage.database import Database
from tests.factories import make_flyer_image

TARGET_ROLES = ["Java Developer", "Java Engineer", "Backend Developer"]
EXCLUDE_KEYWORDS = ["course", "bootcamp", "training"]
MUMBAI_ALIASES = ["Mumbai", "Bombay", "Mumbai Metropolitan Region", "MMR"]
REMOTE_TERMS = ["Remote", "Fully Remote", "Remote - India", "WFH"]


def _build_fixture(tmp_path):
    flyer = make_flyer_image(
        tmp_path / "flyer.png",
        ["URGENT HIRING", "Java Developer", "Mumbai", "flyer@example.com"],
    )
    posts = [
        {
            "post_url": "https://linkedin.com/posts/1",
            "post_text": "Hiring Java Developer\nLocation: Mumbai\nSend CV to jobs@example.com",
            "author_name": "John Doe",
            "author_linkedin_url": "https://linkedin.com/in/johndoe",
        },
        {
            "post_url": "https://linkedin.com/posts/2",
            "post_text": "Hiring Java Developer\nLocation: Pune\nSend CV to careers@example.com",
            "author_name": "Jane Smith",
        },
        {
            "post_url": "https://linkedin.com/posts/3",
            "post_text": "Hiring Java Developer\nLocation: Bangalore\nRemote - India\nhr@example.com",
            "author_name": "Priya Rao",
        },
        {
            "post_url": "https://linkedin.com/posts/4",
            "post_text": "Learn Java Developer skills. Join our Java Bootcamp today!",
            "author_name": "Course Provider",
        },
        {
            # Same email as post 1 -> should dedupe against it.
            "post_url": "https://linkedin.com/posts/5",
            "post_text": "Hiring Java Developer\nLocation: Mumbai\nSend CV to jobs@example.com",
            "author_name": "Someone Else",
        },
        {
            "post_url": "https://linkedin.com/posts/6",
            "post_text": "",
            "image_urls": [str(flyer)],
            "author_name": "Flyer Poster",
        },
    ]
    fixture_path = tmp_path / "posts.json"
    fixture_path.write_text(json.dumps(posts))
    return fixture_path


def _config():
    return Config(
        target_roles=TARGET_ROLES,
        exclude_keywords=EXCLUDE_KEYWORDS,
        accepted_location_aliases=MUMBAI_ALIASES,
        accepted_remote_terms=REMOTE_TERMS,
        dry_run=True,
    )


def test_full_pipeline_run(tmp_path):
    fixture_path = _build_fixture(tmp_path)
    db = Database(tmp_path / "test.db")

    summary = run_pipeline(
        _config(),
        db,
        discover_fn=discover_posts,
        discover_arg=fixture_path,
        checkpoint_path=tmp_path / "checkpoint.json",
        logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
    )

    assert summary.posts_discovered == 6
    assert summary.posts_processed == 6
    assert summary.posts_failed == 0

    assert summary.mumbai_matches == 3  # posts 1, 5 (dup, but still Mumbai), and 6
    assert summary.remote_matches == 1  # post 3
    assert summary.rejected_locations == 1  # post 2
    assert summary.excluded == 1  # post 4
    assert summary.duplicates == 1  # post 5 (dupes post 1's email)
    assert summary.records_requiring_review == 3  # posts 1, 3, 6

    assert summary.images_found == 1
    assert summary.images_processed == 1
    assert summary.ocr_successes == 1
    assert summary.ocr_failures == 0

    # Exports: all 6 records included regardless of status (diagnostics).
    wb = openpyxl.load_workbook(tmp_path / "data" / "linkedin_recruiters.xlsx")
    assert wb.active.max_row == 7  # header + 6 records

    with open(tmp_path / "data" / "linkedin_recruiters.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 6

    text_content = (tmp_path / "data" / "linkedin_recruiters.txt").read_text(encoding="utf-8")
    assert text_content.count("LinkedIn Post:") == 6

    # dry_run=True -> approved_recruiters.xlsx must NOT be written, even
    # though nothing is APPROVED yet anyway (no review has happened).
    assert not (tmp_path / "data" / "approved_recruiters.xlsx").exists()

    # Logs
    with open(tmp_path / "logs" / "success.csv", newline="", encoding="utf-8") as f:
        assert len(list(csv.DictReader(f))) == 3
    with open(tmp_path / "logs" / "skipped.csv", newline="", encoding="utf-8") as f:
        assert len(list(csv.DictReader(f))) == 3  # posts 2, 4, 5
    assert (tmp_path / "logs" / "discovery.log").exists()
    assert not (tmp_path / "logs" / "ocr-failures.csv").exists()  # no OCR failures occurred

    # Checkpoint
    checkpoint = read_checkpoint(tmp_path / "checkpoint.json")
    assert checkpoint["posts_seen"] == 6
    assert checkpoint["last_post_url"] == "https://linkedin.com/posts/6"

    db.close()


def test_resume_skips_already_processed_posts(tmp_path):
    fixture_path = _build_fixture(tmp_path)
    db = Database(tmp_path / "test.db")

    run_pipeline(
        _config(),
        db,
        discover_fn=discover_posts,
        discover_arg=fixture_path,
        checkpoint_path=tmp_path / "checkpoint.json",
        logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
    )

    # Second run against the same fixture and same DB - everything was
    # already seen, so nothing new should be processed this time.
    second_summary = run_pipeline(
        _config(),
        db,
        discover_fn=discover_posts,
        discover_arg=fixture_path,
        checkpoint_path=tmp_path / "checkpoint.json",
        logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
    )

    assert second_summary.posts_discovered == 6  # discovery doesn't know what's "new"
    assert second_summary.posts_processed == 0  # but nothing gets reprocessed
    db.close()


def test_non_dry_run_exports_approved_recruiters_file_when_empty(tmp_path):
    fixture_path = _build_fixture(tmp_path)
    db = Database(tmp_path / "test.db")

    run_pipeline(
        _config(),
        db,
        discover_fn=discover_posts,
        discover_arg=fixture_path,
        checkpoint_path=tmp_path / "checkpoint.json",
        logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
        dry_run=False,
    )

    # File is written (nothing is APPROVED yet without a review pass, but the
    # file itself - header only - should exist since dry_run is False).
    path = tmp_path / "data" / "approved_recruiters.xlsx"
    assert path.exists()
    wb = openpyxl.load_workbook(path)
    assert wb.active.max_row == 1  # header only, nothing approved yet
    db.close()


def test_scroll_settings_from_config_reach_the_discover_function(tmp_path):
    """config.json's max_scrolls / scroll_pause_seconds must actually arrive
    at the discover function - otherwise editing them does nothing."""
    received = {}

    def recording_discover(arg, search_terms, max_posts, max_scrolls=None, scroll_pause_seconds=None):
        received.update(
            arg=arg,
            search_terms=search_terms,
            max_posts=max_posts,
            max_scrolls=max_scrolls,
            scroll_pause_seconds=scroll_pause_seconds,
        )
        return []

    config = Config(
        search_terms=["Java Developer"],
        max_posts=250,
        max_scrolls=40,
        scroll_pause_seconds=2.0,
    )
    db = Database(tmp_path / "test.db")
    run_pipeline(
        config,
        db,
        discover_fn=recording_discover,
        discover_arg="whatever",
        checkpoint_path=tmp_path / "checkpoint.json",
        logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
    )

    assert received["max_posts"] == 250
    assert received["max_scrolls"] == 40
    assert received["scroll_pause_seconds"] == 2.0
    db.close()


def test_config_loads_new_scroll_fields_from_json(tmp_path):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({"max_scrolls": 60, "scroll_pause_seconds": 2.5}))
    config = Config.load(path)
    assert config.max_scrolls == 60
    assert config.scroll_pause_seconds == 2.5


def test_config_scroll_defaults_match_previous_hardcoded_values():
    config = Config()
    assert config.max_scrolls == 15
    assert config.scroll_pause_seconds == 1.5


def test_should_stop_ends_run_early(tmp_path):
    fixture_path = _build_fixture(tmp_path)
    db = Database(tmp_path / "stop.db")
    calls = {"n": 0}

    def stop_after_two():
        calls["n"] += 1
        return calls["n"] > 2

    summary = run_pipeline(
        _config(), db, discover_fn=discover_posts, discover_arg=fixture_path,
        checkpoint_path=tmp_path / "c.json", logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data", should_stop=stop_after_two,
    )
    assert summary.posts_discovered == 6
    assert summary.posts_processed == 2


def test_summary_counts_review_records_without_email(tmp_path):
    fixture_path = _build_fixture(tmp_path)
    db = Database(tmp_path / "noemail.db")
    summary = run_pipeline(
        _config(), db, discover_fn=discover_posts, discover_arg=fixture_path,
        checkpoint_path=tmp_path / "c.json", logs_dir=tmp_path / "logs",
        data_dir=tmp_path / "data",
    )
    assert 0 <= summary.review_without_email <= summary.records_requiring_review
    out = []
    summary.print_report(out.append)
    assert any("no email" in line for line in out)
